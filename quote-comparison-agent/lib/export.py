"""Export comparison summary to Markdown, DOCX, and simple PDF."""

from __future__ import annotations

import io
from datetime import datetime, timezone

from lib.schema.models import ComparisonSummary


def to_markdown(summary: ComparisonSummary, session_id: str) -> str:
    h = summary.header
    lines = [
        "# INTERNAL — Sales team review required",
        f"Session: {session_id} · Generated: {h.run_date or datetime.now(timezone.utc).isoformat()}",
        "",
        "## 1. Header",
        f"- Henley plan: {h.henley_plan}",
        f"- Competitor: {h.competitor_builder} — {h.competitor_plan}",
        f"- Region: {h.region} · Rep: {h.sales_rep}",
        "",
        "## 2. Headline Snapshot",
        summary.headline_snapshot,
        "",
        "## 3. Design Differences",
    ]
    for d in summary.design_differences:
        lines.append(f"- {d.get('point', d)}")

    lines += ["", "## 4. Inclusion Differences", "| Category | Henley | Competitor |", "|---|---|---|"]
    for row in summary.inclusion_differences:
        lines.append(f"| {row.category} | {row.henley} | {row.competitor} |")

    lines += ["", "## 5. Henley Value Advantage"]
    for v in summary.henley_value_advantage:
        val = f"${v.value:,.0f}" if v.value is not None else "TBC"
        lines.append(f"- {v.point} — {val} ({v.source_ref})")

    lines += ["", "## 8. Flagged Items"]
    for f in summary.flagged_items:
        lines.append(f"- **{f.get('itemName')}** [{f.get('confidence')}]: {f.get('reason')}")

    lines += ["", "## 9. Open Questions"]
    for q in summary.open_questions:
        lines.append(f"- {q}")

    return "\n".join(lines)


def to_docx_bytes(summary: ComparisonSummary, session_id: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading("INTERNAL — Sales team review required", 0)
    doc.add_paragraph(f"Session {session_id}")
    doc.add_heading("Headline Snapshot", level=1)
    doc.add_paragraph(summary.headline_snapshot)

    doc.add_heading("Inclusion Differences", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Category"
    hdr[1].text = "Henley"
    hdr[2].text = "Competitor"
    for row in summary.inclusion_differences:
        cells = table.add_row().cells
        cells[0].text = row.category
        cells[1].text = row.henley
        cells[2].text = row.competitor

    doc.add_heading("Henley Value Advantage", level=1)
    for v in summary.henley_value_advantage:
        p = doc.add_paragraph(style="List Bullet")
        val = f"${v.value:,.0f}" if v.value is not None else "TBC"
        p.add_run(f"{v.point} — {val} ({v.source_ref})")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf_bytes(summary: ComparisonSummary, session_id: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 14)
    c.drawString(2 * cm, h - 2 * cm, "INTERNAL — Sales team review required")
    c.setFont("Helvetica", 9)
    c.drawString(2 * cm, h - 2.6 * cm, f"Session {session_id}")
    c.setFont("Helvetica-Bold", 11)
    c.drawString(2 * cm, h - 3.5 * cm, "Headline Snapshot")
    c.setFont("Helvetica", 10)
    text = c.beginText(2 * cm, h - 4.2 * cm)
    text.setLeading(14)
    for line in _wrap(summary.headline_snapshot, 90):
        text.textLine(line)
    c.drawText(text)
    c.setFont("Helvetica", 8)
    c.drawString(2 * cm, 1.5 * cm, "Fusion5 for Henley Homes — INTERNAL watermark")
    c.showPage()
    c.save()
    return buf.getvalue()


def _wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines, cur = [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            lines.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    if cur:
        lines.append(cur)
    return lines
